"""#147: POST /api/preview, GET /preview and GET /api/export, through a real server on 127.0.0.1 (server.py, publish_flow.py).
Preview makes the page as publish does, in memory, with no publication and no file; Export streams the project's workbook as a
download and leaves no temporary file behind. Reuses the rig of test_publish_route.

Run: python3 dev/run_tests.py test_preview_export_route
"""
import http.client
import json
import tempfile
import threading
import time
import unittest
import urllib.parse
from pathlib import Path
from unittest import mock

from _paths import SRC  # noqa: F401
from openpyxl import load_workbook
from clark import excel, excel_import, platform, publish, publish_flow, server as sv
import test_publish_route as base


def get(rig, path, token=None, host=None, headers=None):
    """(status, headers, raw body) for a GET; `token` False sends none."""
    conn = http.client.HTTPConnection(sv.HOST, rig.httpd.server_port, timeout=20)
    try:
        h = {"Host": host or f"127.0.0.1:{rig.httpd.server_port}", **(headers or {})}
        if token is not False:
            h[sv.TOKEN_HEADER] = token or sv.TOKEN
        conn.request("GET", path, None, h)
        r = conn.getresponse()
        return r.status, r, r.read()
    finally:
        conn.close()


class Case(unittest.TestCase):
    def setUp(self):
        self.rig = base.Rig(self)

    def preview(self, **kw):
        return self.rig.call("POST", "/api/preview", {"project": "PUB"}, **kw)


class Preview(Case):
    def test_it_answers_an_address_and_the_address_serves_the_publish_bytes_once(self):
        status, body = self.preview()
        self.assertEqual((status, body["ok"]), (200, True))
        self.assertIn("not published", body["message"])
        self.assertIn("version kept on this computer, not the client's copy", body["message"])
        s2, r, page = get(self.rig, body["url"], token=False)       # a new tab sends no token
        self.assertEqual(s2, 200)
        self.assertTrue(r.getheader("Content-Type").startswith("text/html"))
        self.assertIn("Content-Security-Policy", dict(r.getheaders()))
        want = publish.build(self.rig.store, "PUB", base.DAY, base.EMPTY_BRANDING)["body"]
        self.assertEqual(page, want)
        self.assertEqual(get(self.rig, body["url"], token=False)[0], 404)       # one visit

    def test_it_logs_no_publication_and_writes_no_file(self):
        before = self.rig.everything_written()
        status, body = self.preview()
        get(self.rig, body["url"], token=False)
        self.assertEqual(self.rig.rows(), [])
        self.assertEqual(self.rig.everything_written(), before)
        self.assertEqual(self.rig.saves, [])                      # no Save window either

    def test_it_needs_the_token(self):
        self.assertEqual(self.preview(token=False)[0], 403)
        self.assertEqual(self.preview(token="wrong")[0], 403)

    def test_a_foreign_origin_is_refused_for_the_post_and_for_the_address(self):
        _, body = self.preview()
        status, r, raw = get(self.rig, body["url"], token=False, headers={"Origin": "http://evil.example"})
        self.assertEqual(status, 403)
        self.assertEqual(json.loads(raw)["code"], "not_allowed")
        conn = http.client.HTTPConnection(sv.HOST, self.rig.httpd.server_port, timeout=20)
        data = json.dumps({"project": "PUB"}).encode()
        conn.request("POST", "/api/preview", data, {"Host": f"127.0.0.1:{self.rig.httpd.server_port}", sv.TOKEN_HEADER: sv.TOKEN,
                                                    "Origin": "http://evil.example", "Content-Type": "application/json",
                                                    "Content-Length": str(len(data))})
        self.assertEqual(conn.getresponse().status, 403)
        conn.close()
        self.assertEqual(get(self.rig, body["url"], token=False, headers={"Sec-Fetch-Site": "cross-site"})[0], 403)

    def test_a_guess_or_a_wrong_host_gets_nothing(self):
        self.assertEqual(get(self.rig, "/preview?id=guess", token=False)[0], 404)
        self.assertEqual(get(self.rig, "/preview", token=False)[0], 404)
        _, body = self.preview()
        self.assertEqual(get(self.rig, body["url"], token=False, host="evil.example")[0], 403)

    def test_a_bad_or_missing_project_is_said_in_words(self):
        self.assertEqual(self.rig.call("POST", "/api/preview", {"project": "x"})[0], 400)
        status, body = self.rig.call("POST", "/api/preview", {"project": "ZZZ"})
        self.assertEqual(status, 404)
        self.assertEqual(body["message"], "There is no project with that code.")

    def test_a_page_that_cannot_be_built_says_what_to_do(self):
        with mock.patch.object(publish, "build", side_effect=publish.PublishError("An archived project can't be published. Restore it first.")):
            status, body = self.preview()
        self.assertEqual(status, 409)
        self.assertEqual(body["message"], "The preview could not be made. An archived project can't be published. Restore it first.")

    def test_a_preview_not_opened_in_time_is_gone(self):
        url = self.preview()[1]["url"]
        with mock.patch("clark.publish_flow.time.monotonic", return_value=time.monotonic() + publish_flow.PREVIEW_SECONDS + 1):
            self.assertEqual(get(self.rig, url, token=False)[0], 404)

    def test_only_a_few_previews_wait(self):
        urls = [self.preview()[1]["url"] for _ in range(6)]
        self.assertEqual(get(self.rig, urls[0], token=False)[0], 404)
        self.assertEqual(get(self.rig, urls[-1], token=False)[0], 200)


class Export(Case):
    def test_it_downloads_the_workbook_named_with_the_code_and_the_day(self):
        status, r, raw = get(self.rig, "/api/export?project=PUB")
        self.assertEqual(status, 200)
        self.assertEqual(r.getheader("Content-Disposition"), 'attachment; filename="PUB-2026-10-28.xlsx"')
        self.assertEqual(r.getheader("Content-Type"), sv.XLSX_TYPE)
        self.assertEqual(r.getheader("Content-Length"), str(len(raw)))
        self.assertEqual(r.getheader("Cache-Control"), "no-store")
        self.assertEqual(r.getheader("X-Content-Type-Options"), "nosniff")
        self.assertEqual(raw[:2], b"PK")
        with tempfile.TemporaryDirectory() as t:
            f = Path(t) / "x.xlsx"
            f.write_bytes(raw)
            self.assertIn("milestones", {k for k, v in excel.read_workbook(f).items() if v})

    def test_it_is_the_same_content_as_the_export_of_the_store(self):
        with tempfile.TemporaryDirectory() as t:
            want = excel.export_project(self.rig.store, "PUB", Path(t) / "w.xlsx", Path(t) / "s.json", on=base.DAY)
            direct = excel.read_workbook(want)
            _, _, raw = get(self.rig, "/api/export?project=PUB")
            (Path(t) / "got.xlsx").write_bytes(raw)
            self.assertEqual(excel.read_workbook(Path(t) / "got.xlsx"), direct)

    def test_nothing_is_left_behind_and_only_the_record_is_written_in_the_data_folder(self):
        before, tmp = self.rig.everything_written(), set(Path(tempfile.gettempdir()).glob("clark-export-*"))
        get(self.rig, "/api/export?project=PUB")
        self.assertEqual(set(Path(tempfile.gettempdir()).glob("clark-export-*")), tmp)
        self.assertEqual([p.name for p in self.rig.data_dir.iterdir() if p.suffix == ".xlsx"], [])
        self.assertTrue((self.rig.data_dir / excel.STATE_FILE).exists())
        self.assertEqual(self.rig.rows(), [])

    def test_the_download_is_recorded_so_it_imports_as_recorded_from_any_folder_while_unchanged(self):
        _, _, raw = get(self.rig, "/api/export?project=PUB")
        state = self.rig.data_dir / excel.STATE_FILE
        for folder in ("Downloads", "Documents"):                  # wherever the browser or the person put the downloaded bytes
            (Path(self.rig.root) / folder).mkdir(exist_ok=True)
            mine = Path(self.rig.root) / folder / "PUB-2026-10-28.xlsx"
            mine.write_bytes(raw)
            plan = excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state)
            self.assertFalse(plan.unrecorded, folder)
        # a file whose bytes differ is not the export: nothing says what it was made from
        other = Path(self.rig.root) / "elsewhere.xlsx"
        other.write_bytes(raw + b"x")
        self.assertTrue(excel_import.read_workbook(self.rig.store, other, "PUB", state_path=state, allow_unrecorded=False).unrecorded)
        self.assertEqual(excel.load_state(state, other), {})

    def test_known_limit_until_169_a_workbook_saved_again_with_no_edit_is_unrecorded(self):
        """The record is found only while the file is exactly what clark wrote. A re-save (openpyxl, no edit) changes the bytes, so the
        import calls it unrecorded. This pins the limit; #169 (an export id inside the workbook) is what lifts it, and then this changes."""
        _, _, raw = get(self.rig, "/api/export?project=PUB")
        state, mine = self.rig.data_dir / excel.STATE_FILE, Path(self.rig.root) / "PUB-resaved.xlsx"
        mine.write_bytes(raw)
        self.assertFalse(excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state).unrecorded)
        wb = load_workbook(mine)
        wb.save(mine)                                                                        # saved again, nothing edited
        self.assertNotEqual(mine.read_bytes(), raw)
        self.assertTrue(excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state).unrecorded)

    def test_records_kept_by_bytes_older_than_thirty_days_are_dropped_on_the_next_export(self):
        get(self.rig, "/api/export?project=PUB")
        state = self.rig.data_dir / excel.STATE_FILE
        whole = json.loads(state.read_text(encoding="utf-8"))
        whole["workbooks"]["sha256:" + "0" * 64] = {"exported_at": "2026-08-01 09:00:00", "hash": "0" * 64, "rows": {}}
        whole["workbooks"]["/some/path.xlsx"] = {"exported_at": "2026-08-01 09:00:00", "hash": "1" * 64, "rows": {}}
        state.write_text(json.dumps(whole), encoding="utf-8")
        get(self.rig, "/api/export?project=PUB")
        kept = json.loads(state.read_text(encoding="utf-8"))["workbooks"]
        self.assertNotIn("sha256:" + "0" * 64, kept)
        self.assertIn("/some/path.xlsx", kept)
        self.assertEqual(len([k for k in kept if k.startswith("sha256:")]), 1)

    def test_a_second_export_on_the_same_day_does_not_replace_the_first_ones_record(self):
        _, _, first = get(self.rig, "/api/export?project=PUB")
        with self.rig.store.step("dashboard", base.USER) as st:
            st.create("actions", {"project": "PUB", "description": "Added between the two exports", "raised": "2026-10-28", "source": "test"})
        _, _, second = get(self.rig, "/api/export?project=PUB")
        self.assertNotEqual(first, second)
        state = self.rig.data_dir / excel.STATE_FILE
        for n, raw in (("one", first), ("two", second)):
            f = Path(self.rig.root) / f"{n}.xlsx"
            f.write_bytes(raw)
            self.assertEqual(excel.load_state(state, f).get("hash"), excel.file_hash(f), n)

    def test_it_needs_the_token_and_refuses_a_foreign_origin(self):
        self.assertEqual(get(self.rig, "/api/export?project=PUB", token=False)[0], 403)
        self.assertEqual(get(self.rig, "/api/export?project=PUB", headers={"Origin": "http://evil.example"})[0], 403)
        self.assertEqual(get(self.rig, "/api/export?project=PUB", headers={"Sec-Fetch-Site": "cross-site"})[0], 403)

    def test_a_bad_or_missing_project_is_said_in_words(self):
        self.assertEqual(get(self.rig, "/api/export?project=x")[0], 400)
        s, _, raw = get(self.rig, "/api/export?project=ZZZ")
        self.assertEqual((s, json.loads(raw)["message"]), (404, "There is no project with that code."))

    def test_a_refusal_of_any_kind_is_said_as_not_made_and_cleaned_up(self):
        tmp = set(Path(tempfile.gettempdir()).glob("clark-export-*"))
        for e in (excel.Refused("export.xlsx is open in Excel (/some/path)."), excel.TooBig("too big /some/path")):
            with mock.patch.object(excel, "export_project", side_effect=e):
                s, _, raw = get(self.rig, "/api/export?project=PUB")
            self.assertEqual((s, json.loads(raw)["message"]), (409, "The workbook could not be made. Nothing was changed. Try again."))
            self.assertNotIn("/some/path", raw.decode())
        self.assertEqual(set(Path(tempfile.gettempdir()).glob("clark-export-*")), tmp)

    def test_the_cover_names_the_exporter_and_sends_changes_back_to_whoever_sent_it(self):
        _, _, raw = get(self.rig, "/api/export?project=PUB")
        f = Path(self.rig.root) / "cover.xlsx"
        f.write_bytes(raw)
        text = " | ".join(str(c.value) for row in load_workbook(f)["Cover"].iter_rows() for c in row if c.value)
        self.assertIn("by " + base.USER, text)        # the name publish records (the rig's by_user)
        self.assertIn("send the file back to the person who sent it to you", text)

    def test_a_logo_warning_travels_with_the_download(self):
        def fake(store, code, path, state, on=None, warnings=None, **kw):
            warnings.append("The logo was not used: it is too big.")
            return real(store, code, path, state, on=on, warnings=warnings, **kw)
        real = excel.export_project
        with mock.patch.object(excel, "export_project", side_effect=fake):
            s, r, raw = get(self.rig, "/api/export?project=PUB")
        self.assertEqual((s, raw[:2]), (200, b"PK"))
        self.assertEqual(urllib.parse.unquote(r.getheader("X-Export-Warning")), "The logo was not used: it is too big.")

    def test_a_fault_while_making_it_is_plain_words(self):
        with mock.patch.object(excel, "export_project", side_effect=OSError("disk /secret full")):
            s, _, raw = get(self.rig, "/api/export?project=PUB")
        self.assertEqual((s, json.loads(raw)["message"]), (409, "The workbook could not be made. Nothing was changed. Try again."))
        self.assertNotIn("secret", raw.decode())


class ExportRecord(Case):
    """The record each export keeps beside clark.db (excel.STATE_FILE), as the comments in excel.py and excel_import.py say."""

    def download(self, folder="Downloads", name="PUB-2026-10-28.xlsx"):
        _, _, raw = get(self.rig, "/api/export?project=PUB")
        (Path(self.rig.root) / folder).mkdir(exist_ok=True)
        mine = Path(self.rig.root) / folder / name
        mine.write_bytes(raw)
        return mine, raw

    def edit_milestone(self, path, words):
        wb = load_workbook(path)
        ws = wb["Milestones"]
        cell = [c for row in ws.iter_rows() for c in row if str(c.value or "").startswith("Pilot live")][0]
        cell.value = words
        wb.save(path)

    def test_two_exports_at_once_both_keep_their_record(self):
        """Export A is held at its write of the state file while a milestone is renamed and export B runs. Without one lock around
        the read, change and write, B's record is written and then A's write (of the file as A read it) throws it away."""
        state = self.rig.data_dir / excel.STATE_FILE
        real, at_write, go, b_wrote, first = platform.write_file, threading.Event(), threading.Event(), threading.Event(), []

        def held(path, *a, **kw):
            if Path(path) == state:
                if not first:
                    first.append(threading.current_thread())
                    at_write.set()
                    go.wait(10)                           # A waits here, holding what it read
                elif threading.current_thread() is not first[0]:
                    b_wrote.set()
            return real(path, *a, **kw)
        got = {}

        def export(n):
            got[n] = get(self.rig, "/api/export?project=PUB")
        with mock.patch.object(platform, "write_file", held):
            a = threading.Thread(target=export, args=("A",))
            a.start()
            self.assertTrue(at_write.wait(10))
            with self.rig.store.step("dashboard", base.USER) as st:
                st.update("milestones", "PUB-M-001", {"name": "Pilot live, renamed between the exports"})
            b = threading.Thread(target=export, args=("B",))
            b.start()
            b_wrote.wait(2)                               # without the lock B writes now; with it B waits for A
            go.set()
            a.join(20)
            b.join(20)
        self.assertEqual((got["A"][0], got["B"][0]), (200, 200))
        self.assertNotEqual(got["A"][2], got["B"][2])
        for n in ("A", "B"):
            f = Path(self.rig.root) / f"{n}.xlsx"
            f.write_bytes(got[n][2])
            self.assertEqual(excel.load_state(state, f).get("hash"), excel.file_hash(f), f"export {n} has no record")
            self.assertFalse(excel_import.read_workbook(self.rig.store, f, "PUB", state_path=state).unrecorded, n)

    def test_known_limit_until_169_an_edited_workbook_is_unrecorded_and_stays_so_after_a_confirmed_import(self):
        """What load_state's comment says: an edit loses the record, and a confirmed import leaves a record under the path with no
        export in it, so the next edit of the same file is unrecorded again. #169 lifts this, and then this changes."""
        state = self.rig.data_dir / excel.STATE_FILE
        mine, raw = self.download()
        self.edit_milestone(mine, "Pilot live, moved in Excel")
        plan = excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state)
        self.assertTrue(plan.unrecorded)
        self.assertIn("it is recorded only while it is exactly what clark wrote", " ".join(plan.notes))
        self.assertEqual(excel.load_state(state, mine), {})
        plan.confirm_unrecorded()
        plan.accept(*[c.n for c in plan.changes if c.kind == "changed"])
        self.assertTrue(excel_import.apply_import(self.rig.store, plan, base.USER, state_path=state).applied)
        left = excel.load_state(state, mine)                          # the import's record, under the path
        self.assertEqual(left.get("imported_hash"), excel.file_hash(mine))
        self.assertFalse(left.get("hash") or left.get("rows"))
        self.edit_milestone(mine, "Pilot live, moved again")
        self.assertTrue(excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state).unrecorded)

    def test_the_same_bytes_imported_into_another_project_have_no_record(self):
        with self.rig.store.step("dashboard", base.USER) as st:
            st.create("projects", {"code": "OTH", "name": "Other Programme", "health": "On track"})
        state = self.rig.data_dir / excel.STATE_FILE
        mine, _ = self.download()
        self.assertEqual(excel.load_state(state, mine).get("project"), "PUB")
        self.assertFalse(excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state).unrecorded)
        self.assertTrue(excel_import.read_workbook(self.rig.store, mine, "OTH", state_path=state).unrecorded)

    def test_a_record_kept_by_bytes_is_aged_by_its_export_and_the_path_record_an_import_left_is_kept(self):
        state = self.rig.data_dir / excel.STATE_FILE
        mine, _ = self.download()
        excel.mark_imported(mine, state, project="PUB", step={"step": 1, "db_id": "x", "at": "2026-10-28 09:00:00"})
        whole = json.loads(state.read_text(encoding="utf-8"))
        self.assertEqual(sum(k.startswith("sha256:") for k in whole["workbooks"]), 1)    # the import did not touch it...
        self.assertIn(str(mine.resolve()), whole["workbooks"])                         # ...it wrote under the path
        excel.prune_hash_records(state, "2026-12-15")
        kept = json.loads(state.read_text(encoding="utf-8"))["workbooks"]
        self.assertEqual([k for k in kept if k.startswith("sha256:")], [])
        self.assertEqual(kept[str(mine.resolve())]["imported_hash"], excel.file_hash(mine))
        # the import's record under the path holds the export too, so ageing the bytes-keyed one by its export date lost nothing
        self.assertEqual(kept[str(mine.resolve())]["hash"], excel.file_hash(mine))
        self.assertTrue(kept[str(mine.resolve())]["rows"])
        self.assertFalse(excel_import.read_workbook(self.rig.store, mine, "PUB", state_path=state).unrecorded)


class ExportSnapshotBound(Case):
    """#181: an Export click no longer grows the record file without limit, and an archived project keeps no snapshot."""

    def books(self):
        state = self.rig.data_dir / excel.STATE_FILE
        return json.loads(state.read_text(encoding="utf-8"))["workbooks"] if state.exists() else {}

    def hashed(self):
        return [k for k in self.books() if k.startswith("sha256:")]

    def change(self, n):
        with self.rig.store.step("dashboard", base.USER) as st:
            st.create("actions", {"project": "PUB", "description": f"Change {n}", "raised": "2026-10-28", "source": "test"})

    def test_fifty_clicks_leave_one_record_and_the_record_file_is_not_written_again(self):
        get(self.rig, "/api/export?project=PUB")
        state = self.rig.data_dir / excel.STATE_FILE
        first = state.read_bytes()
        with mock.patch.object(platform, "write_file", wraps=platform.write_file) as write:
            for _ in range(49):
                get(self.rig, "/api/export?project=PUB")
            wrote = [c.args[0] for c in write.call_args_list if Path(c.args[0]) == state]
        self.assertEqual(len(self.hashed()), 1)
        self.assertEqual(state.read_bytes(), first)
        self.assertEqual(wrote, [], "the same record is never written to the record file again")

    def test_exports_that_change_keep_only_the_newest_ten_for_the_project(self):
        raws = []
        for n in range(excel.KEEP_NEWEST + 3):
            self.change(n)
            raws.append(get(self.rig, "/api/export?project=PUB")[2])
        self.assertEqual(len(self.hashed()), excel.KEEP_NEWEST)
        state = self.rig.data_dir / excel.STATE_FILE
        for i, raw in enumerate(raws):
            f = Path(self.rig.root) / f"e{i}.xlsx"
            f.write_bytes(raw)
            self.assertEqual(bool(excel.load_state(state, f)), i >= len(raws) - excel.KEEP_NEWEST, i)

    def test_an_archived_project_keeps_no_snapshot_at_all(self):
        with self.rig.store.step("dashboard", base.USER) as st:
            st.archive("projects", "PUB")
        get(self.rig, "/api/export?project=PUB")
        self.assertEqual(self.hashed(), [])


if __name__ == "__main__":
    unittest.main()
