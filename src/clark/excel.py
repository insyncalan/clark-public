"""The Excel export (#37; BRD FR-70): one project's actions, RAID, decisions, milestones and use cases as a workbook.

Copied from creaite v2.6.9 (db5adb2), system/scripts/excel_sync.py: the sheet layout, the cell conversions (to_cell, from_cell,
cell_text), the drop-downs and locked columns, and the guard against overwriting a workbook with edits not yet imported (a
record of the SHA-256 of the file each export wrote, and of the last one imported). Each copied function names its original.
What is cut: the Projects, Workstreams, Promises, Links, Updates and Review sheets, links in plain text (#121; clark has no
linktext), the import itself, and creaite's paths and debug log. What differs is marked "clark:".

Formatting (#85a, #85c): the header in the brand navy, the ID and the item's title frozen, widths and wrapping by kind, and
conditional formatting from excel_colours.py (clark's own; creaite has none): a rule per value of Status, Priority, Impact,
Likelihood and State, so a value changed in Excel recolours itself, and a formula rule that makes the date of an open item past
its date red. The rules are fixed literals, never data; no cell holds a formula, and colour never changes a cell's value. Every
fixed-value column has a drop-down that refuses anything else; the project's own lists (workstreams, milestones, use cases) are
on a hidden "Lists" sheet that the import never reads. Owner stays free text. "Updated" stays the stored UTC text,
because the import compares it to spot a clash; it is not shown as a local time. Other stored times in the file (the state
file's exported_at, the workbook's created and modified times) are UTC too; a time shown to a person goes through
clark.timefmt.local_text (exported_text, for the cover).

The workbook is never the record: the store is. Dates are real Excel dates shown DD-MM-YYYY; they are stored YYYY-MM-DD.

`export()` takes the rows themselves ({table: [row dict, ...]}) and needs no database; `export_project()` reads them from the
store (#31) with Store.read. openpyxl is the copy in vendor/; whoever starts clark puts vendor/ on sys.path, as
dev/run_tests.py does.

INTERNAL ONLY: the Milestones sheet carries each milestone's Internal note, and Sensitive and client-hidden rows are exported
too. Never send this workbook to a client. It says so in the file itself (the document description, and a note on Milestones A1).

Text is neutralised on the way out (escape_text) and restored on the way in (from_cell), so a value such as =HYPERLINK(...)
is only ever text and is never run. The file is byte-for-byte the same for the same rows and the same day (see `moment`).
"""
import hashlib
import io
import json
import re
import threading
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.xml.functions import tostring
from openpyxl.worksheet.datavalidation import DataValidation

from clark import excel_colours as colours
from clark import excel_cover
from clark import excel_picture
from clark import platform
from clark.milestones import STATES, milestone_state
from clark.timefmt import local_text, parse

EXTRA_ROWS = 500          # rows below the data that keep drop-downs and formats for new entries
DATE_FORMAT = "dd-mm-yyyy"
STATE_FILE = "excel-export-state.json"   # beside clark.db: what each export wrote, by workbook
# One read-change-write of a state file at a time (#147): the server answers each request on its own thread, and two exports
# (or an export and an import) that each read the whole file, add their record and write it back would otherwise lose one.
# Re-entrant, so mark_imported can hold it across its load_state and save_state. It covers one running clark only, not two
# clarks started on the same data folder.
STATE_LOCK = threading.RLock()

# Sheet layout (creaite excel_sync.SHEETS): (sheet name, table, [(column label, field)], locked fields).
# clark: + Workstream and Movement on milestones, + Use case on actions, RAID and decisions, + the Use cases sheet; Original
# and Shown state are worked out, not columns (see DERIVED); no Labels on milestones or use cases, as in creaite.
SHEETS = [
    ("Milestones", "milestones", [
        ("ID", "id"), ("Milestone", "name"), ("Project", "project"), ("Description", "description"),
        ("Workstream", "workstream"), ("Original", "original_date"), ("Agreed", "baseline_date"),
        ("Forecast", "forecast_date"), ("Status", "status"), ("Order", "sort_order"), ("% complete", "percent_complete"),
        ("Owner", "owner"), ("Shown state", "shown_state"), ("State", "state"), ("State set by you", "state_set"),
        ("State reason", "state_reason"), ("Movement", "movement"), ("Achieved on", "actual_date"),
        ("Client-visible", "client_visible"), ("Client wording", "client_message"), ("Internal note", "internal_note"),
        ("Source", "source"), ("Updated", "updated_at")],
     {"id", "original_date", "shown_state", "source", "updated_at"}),
    ("Actions", "actions", [
        ("ID", "id"), ("Description", "description"), ("Project", "project"), ("Workstream", "workstream"),
        ("Owner", "owner"), ("Priority", "priority"), ("Due", "due_date"), ("Status", "status"),
        ("Raised", "raised"), ("Closed", "closed"), ("Latest update", "latest_update"),
        ("Client-visible", "client_visible"), ("Client wording", "client_message"), ("Milestone", "milestone"),
        ("Use case", "use_case"), ("Labels", "labels"), ("Source", "source"), ("Updated", "updated_at")],
     {"id", "source", "updated_at"}),
    ("Risks and issues", "raid", [
        ("ID", "id"), ("Description", "description"), ("Project", "project"), ("Type", "type"),
        ("Workstream", "workstream"), ("Owner", "owner"), ("Response action", "response_action"),
        ("Due", "due_date"), ("Impact", "impact"), ("Likelihood", "likelihood"), ("Priority", "priority"),
        ("Status", "status"), ("Raised", "raised"), ("Closed", "closed"), ("Client-visible", "client_visible"),
        ("Client wording", "client_message"), ("Milestone", "milestone"), ("Use case", "use_case"), ("Labels", "labels"),
        ("Source", "source"), ("Updated", "updated_at")],
     {"id", "priority", "source", "updated_at"}),
    ("Decisions", "decisions", [
        ("ID", "id"), ("Decision", "decision"), ("Project", "project"), ("Status", "status"), ("Made by", "made_by"),
        ("Needed by", "needed_by"), ("Date", "date"), ("Reasoning", "reasoning"), ("Options", "options"),
        ("Revisit", "revisit"), ("Replaces", "replaces"), ("Client-visible", "client_visible"),
        ("Client wording", "client_message"), ("Use case", "use_case"), ("Labels", "labels"), ("Source", "source"),
        ("Updated", "updated_at")],
     {"id", "source", "updated_at"}),
    ("Use cases", "use_cases", [                                       # clark's own table, in creaite's style
        ("ID", "id"), ("Use case", "title"), ("Project", "project"), ("Shown as", "ref"), ("Status", "status"),
        ("Headline", "headline"), ("Where it stands", "standing"), ("Order", "sort_order"),
        ("Client-visible", "client_visible"), ("Source", "source"), ("Updated", "updated_at")],
     {"id", "source", "updated_at"}),
]
TABLES = tuple(t for _, t, _, _ in SHEETS)
# The same layout as two lookups, for the import (#38): the sheet a table is written on, and each table's header label -> field.
# Derived from SHEETS, so the import reads exactly the labels the export writes.
SHEET_OF = {table: name for name, table, _, _ in SHEETS}
COLUMN_OF = {table: {label: field for label, field in cols} for _, table, cols, _ in SHEETS}
DATES = {"due_date", "raised", "closed", "date", "baseline_date", "forecast_date", "needed_by", "actual_date",
         "original_date"}
YES_NO = {"client_visible", "state_set"}
LABELLED = ("actions", "raid", "decisions")
DERIVED = {"original_date", "shown_state"}    # clark: worked out for the sheet, never read back as a change
WIDE = {"description", "decision", "client_message", "latest_update", "response_action", "reasoning", "options", "revisit",
        "state_reason", "internal_note", "source", "headline", "standing"}
# clark: a NOT NULL DEFAULT '' text column. A blank cell there reads back as "" (creaite's blank_default, with no template: #24).
EMPTY_TEXT = {("actions", "client_message"), ("raid", "client_message"), ("decisions", "client_message"),
              ("use_cases", "ref"), ("use_cases", "headline"), ("use_cases", "standing")}
CHOICES = {
    ("actions", "status"): ["Open", "Delayed", "Closed", "Cancelled"],
    ("raid", "status"): ["Open", "Closed", "Cancelled"],
    ("milestones", "status"): ["Planned", "Done", "Cancelled"],
    ("milestones", "state"): [s for s in STATES if s != "Achieved"] + ["Achieved"],   # clark: + Forecast
    ("milestones", "movement"): ["hold", "rebaseline", "recover", "escalate"],         # clark: FR-53
    ("decisions", "status"): ["Pending", "Made", "Withdrawn"],
    ("use_cases", "status"): ["On track", "Flag", "High risk"],                          # clark
    "priority": ["Urgent", "High", "Medium", "Low"],
    "type": ["Risk", "Assumption", "Issue", "Dependency"],
    "impact": ["High", "Medium", "Low"],
    "likelihood": ["High", "Medium", "Low"],
}

# clark (#85c): NOT NULL columns, whose drop-down does not allow a blank. The rest (and every other sheet's optional links) do.
NOT_BLANK = {("actions", "priority"), ("actions", "status"), ("raid", "type"), ("raid", "status"), ("milestones", "status"),
             ("decisions", "status"), ("use_cases", "status")}
PEOPLE = {"owner", "made_by"}              # free text until the people table exists (#115)
LISTS_NAME = "Lists"                       # the hidden sheet holding the project's own lists
LIST_ERROR_TITLE = "Not on the list"
LIST_ERROR = "Pick one of the values in the list."

INTERNAL_NOTE = ("INTERNAL: this workbook includes internal notes and Sensitive items. Never share it with a client.")
FIXED_NAME = "clark"                       # creator and last-modified-by: fixed and neutral, so the file never names a person
ZIP_TIME = (1980, 1, 1, 0, 0, 0)           # the earliest time a zip can hold: every member gets it

# Formatting (#85a). Brand colours come from excel_colours (the demo tokens); the layout rules are these.
HEADER = PatternFill("solid", fgColor=colours.HEADER_FILL)
HEADER_FONT = Font(bold=True, color=colours.HEADER_TEXT)
# The item's title column on each sheet: the column frozen beside the ID, so a row stays readable while scrolling across.
TITLE_OF = {"milestones": "name", "actions": "description", "raid": "description", "decisions": "decision",
            "use_cases": "title"}
SHORT_NAMES = {"owner", "made_by", "workstream", "labels"}      # people, streams and labels: a little wider than a status
NUMBERS = {"sort_order", "percent_complete"}
LOCKED = PatternFill("solid", fgColor="EFEFEF")
# The ID column (#85d, BRD FR-70d). There are no cell locks and no sheet protection: Excel refuses to sort a range that holds
# a locked cell, and sorting must stay free. The import is the guard (a changed existing ID is refused for that row; a row with
# no ID is a new item and gets the next free ID). The ID column is shaded (LOCKED fill) and these words are pinned by a test.
# The Cover is unprotected on purpose, for the same reason.
ID_NOTE = ("IDs are given by clark. Don't change or move an ID, and don't sort the ID column on its own: an ID moved to "
           "another row takes that row's details with it when the file is imported. Leave the ID empty on a new row.")
ID_PROMPT_TITLE = "ID"
ID_PROMPT_EXISTING = "Don't change this ID: if it changes, this row's details go to another item when the file is imported."
ID_PROMPT_NEW = "Leave the ID empty for a new item. clark gives it an ID when the file is imported."




class Refused(Exception):
    """The export did not write the workbook; the message says why, in plain words, and names the file."""


class TooBig(Refused):
    """A workbook that is too big to open safely (a zip bomb, or the wrong file); the message says so in plain words."""


# What a workbook may be before clark opens it. A real export of a large project is far below all four. The sizes are those the
# .xlsx (a zip) declares for its members; Python's zip reader never gives back more than a member declares.
MAX_FILE_BYTES = 25 * 1024 * 1024          # the file as saved
MAX_UNPACKED_BYTES = 50 * 1024 * 1024      # all its members, unpacked
MAX_MEMBER_BYTES = 40 * 1024 * 1024        # any one member, unpacked
MAX_MEMBERS = 200                          # members in the zip
MAX_SNAPSHOT_BYTES = 20 * 1024 * 1024      # the per-row record in the state file; a bigger one is left out (the import copes)


def check_workbook_size(path):
    """Raise TooBig, before anything is unpacked, if the file is larger than the limits above (a small file that inflates to
    gigabytes is refused by its unpacked size, not just its ratio). A missing file, or one that is not a zip, passes: opening it
    says what is wrong with it."""
    path = Path(path)
    try:
        size = path.stat().st_size
    except OSError:
        return
    mb = lambda n: f"{n / 1048576:.0f} MB" if n >= 1048576 else f"{n / 1024:.0f} KB"
    if size > MAX_FILE_BYTES:
        raise TooBig(f"{path.name} is {mb(size)}, more than the {mb(MAX_FILE_BYTES)} clark will open. A clark export is far "
                     "smaller: check it is the right file.")
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, OSError):
        return
    total, biggest = sum(i.file_size for i in infos), max((i.file_size for i in infos), default=0)
    if len(infos) > MAX_MEMBERS:
        raise TooBig(f"{path.name} holds {len(infos)} parts, more than the {MAX_MEMBERS} a workbook should have, so it was "
                     "not opened. Check it is the right file.")
    if biggest > MAX_MEMBER_BYTES or total > MAX_UNPACKED_BYTES:
        raise TooBig(f"{path.name} would unpack to about {mb(total)} (the most clark will open is {mb(MAX_UNPACKED_BYTES)}), "
                     "so it was not opened. A clark export is far smaller: check it is the right file.")


# ── Helpers ───────────────────────────────────────────────────────────────────

def clock():
    """The real time now, as UTC without a zone mark (#75). The one place the export reads the clock; tests replace it
    (excel.clock = ...) with a naive UTC datetime. UTC because the workbook's created/modified times are UTC in the file (Excel
    shows them in the user's zone) and the state file's exported_at is a machine value: neither depends on the machine's zone."""
    return datetime.now(timezone.utc).replace(tzinfo=None)


def moment(on=None):
    """The time the export is stamped with, to the second, in UTC (#75): midnight UTC at the start of day `on` (YYYY-MM-DD) when
    it is given, else the clock. It sets the workbook's created and modified times and the state file's exported_at, so the
    same rows with the same `on` always give the same bytes, in any zone."""
    if on:
        return datetime.combine(date.fromisoformat(str(on)), datetime.min.time())
    return clock().replace(microsecond=0)


def day(on=None):
    """The calendar day of the export, as a date: `on` when given, else today on the user's own calendar (the clock's moment
    seen in the local zone, #75). Dates are never shifted by a zone; only 'today' depends on where the user is."""
    if on:
        return date.fromisoformat(str(on))
    return moment().replace(tzinfo=timezone.utc).astimezone().date()


def exported_text(on=None):
    """'Exported on DD-MM-YYYY at HH:MM (UTC+04:00)' in the user's local time (#75), for the cover (#85b). The zone is named by
    its UTC offset, not its name: zone names ('Gulf Standard Time', or just '+04') differ between systems, an offset does not.
    The moment is the export's own (moment(on): the clock, or midnight UTC at the start of `on`), shown through clark.timefmt,
    so it depends on the machine's zone by design: a cover carrying it differs between zones for the same data and clock."""
    stamp = moment(on).strftime("%Y-%m-%d %H:%M:%S")
    shown = local_text(stamp)
    offset = int(parse(stamp).astimezone().utcoffset().total_seconds() // 60)
    sign, offset = ("+" if offset >= 0 else "-"), abs(offset)
    return f"Exported on {shown[:10]} at {shown[11:]} (UTC{sign}{offset // 60:02d}:{offset % 60:02d})"


def file_hash(path):
    """The SHA-256 of the workbook, or None when there is none (creaite excel_sync.file_hash)."""
    path = Path(path)
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


def hash_key(digest):
    """A workbook's entry in the state file when it is known by its bytes, not where it is (#147: the menu's export is
    downloaded by the browser, so clark never learns the path). Found only for a byte-identical file: a re-save or an edit
    changes the bytes and the record is lost, until an export id inside the workbook (#169) replaces this."""
    return "sha256:" + digest


def _key(path):
    """The workbook's entry in the state file: its full path, or a hash_key() as it is. A path matches only that file; a
    hash_key() matches only its exact bytes (see hash_key: lost on a re-save or an edit, until #169)."""
    if isinstance(path, str) and path.startswith("sha256:"):
        return path
    return str(Path(path).resolve())


def load_state(state_path, path):
    """The last export's record for this workbook: its file hash, imported_hash and rows (creaite excel_sync.load_state).
    clark: one state file holds a record per workbook. Found by the path, else by the file's SHA-256 (hash_key). What that means:
    the menu's export is recorded only by its bytes, so the downloaded file is found from any folder while it is exactly what
    clark wrote. Once it is re-saved or edited, this returns {}, or, after an import of it was confirmed, the record that
    import left under its path (mark_imported), which has no export in it ("hash", "rows"). Either way the import treats the
    workbook as unrecorded, until an export id in the workbook (#169). The record names the project it was exported from
    ("project"); the caller checks it (excel_import.read_workbook)."""
    state_path = Path(state_path)
    if not state_path.exists():
        return {}
    books = json.loads(state_path.read_text(encoding="utf-8")).get("workbooks", {})
    found = books.get(_key(path))
    if found is None and not str(path).startswith("sha256:"):
        # #147: no record for this path; an export that was downloaded is recorded by its bytes, so look for these
        digest = file_hash(path)
        found = books.get(hash_key(digest)) if digest else None
    return found if found is not None else {}


KEEP_NEWEST = 10          # the most bytes-keyed ("sha256:") records kept for one project (#181); the oldest go first
KEEP_TOTAL = 30           # and the most the state file holds in all, newest first, plus each project's newest (never cut), so more than 30 live projects keep more than 30 (#181)


def prune_hash_records(state_path, today, keep_days=30, keep_newest=KEEP_NEWEST, live=None, keep_total=KEEP_TOTAL, keep=None):
    """Remove the bytes-keyed ("sha256:") records exported more than keep_days before `today`, so the state file does not grow for
    ever (each holds a copy of every row). Path-keyed records are left alone. #147. A record is aged by its export date only, on
    purpose: an import never writes to the bytes-keyed record (mark_imported writes under the workbook's path, which is kept),
    so the export date is the only date it has. Nor does an import need it afterwards: mark_imported copies the record it found
    by the bytes, with the export in it ("hash", "rows"), under that path, so dropping the bytes-keyed one loses nothing for a
    file that was imported. A file still byte for byte the export, kept unimported for over keep_days, then imports as
    unrecorded: every difference is staged for the person to check, and nothing is lost. #181: also, no project keeps more than
    `keep_newest` of them (the newest by export time), at most `keep_total` stay in all (plus each project's newest, which the total never cuts: with more than `keep_total` live projects the file holds one per project and no more than ten of each), and, when `live` (a set of project codes) is
    given, none of a project outside it (an archived or removed one keeps no snapshot at all). The age rule is not the only brake on
    the file's size; an older export then imports as unrecorded, which stages every difference for checking (fails safe).
    `keep` is the record just written: the caps never cut it, and nor does each project's newest, so the total cap cannot take
    the snapshot a person has just made. An export date after today (a clock set back) ranks at the start of today, below every real export made today."""
    state_path = Path(state_path)
    with STATE_LOCK:
        if not state_path.exists():
            return
        whole = json.loads(state_path.read_text(encoding="utf-8"))
        books, cutoff = whole.get("workbooks", {}), str(date.fromisoformat(str(today)[:10]) - timedelta(days=keep_days))
        now = str(today)[:10]
        old = [k for k, v in books.items() if k.startswith("sha256:") and str(v.get("exported_at") or "9")[:10] < cutoff]
        # #181: none of a project that is not live; no more than `keep_newest` for any one project; no more than `keep_total` in all.
        # Newest by export time, then by the order they were saved in (two exports in one second), so the same records go every time: a hundred exports in one day must not make a state
        # file of a hundred copies of every row.
        left = {}
        for at, (k, v) in enumerate(books.items()):                  # `at` is the order they were saved in: the tie-break for two in one second
            if k.startswith("sha256:") and k not in old:
                if live is not None and v.get("project") not in live:
                    old.append(k)
                else:
                    when = str(v.get("exported_at") or "")
                    left.setdefault(str(v.get("project") or ""), []).append((min(when, now + " 00:00:00"), at, k))     # a future date ranks at the start of today
        must = {keep} if keep else set()
        kept, capped = [], []
        for mine in left.values():
            mine.sort(reverse=True)
            must.add(mine[0][2])                                     # each project's newest is exempt from the total cap
            kept += mine[:max(keep_newest, 1)]
            capped += [k for _, _, k in mine[max(keep_newest, 1):]]
        if keep_total and len(kept) > max(keep_total, 1):
            kept.sort(reverse=True)
            capped += [k for _, _, k in kept[max(keep_total, 1):]]
        old += [k for k in capped if k not in must]
        if old:
            for k in old:
                del books[k]
            platform.write_file(state_path, json.dumps(whole, indent=1, ensure_ascii=False))


def save_state(state_path, path, record):
    """Write this workbook's export record into the state file (creaite excel_sync.save_state), safely (platform.write_file).
    clark: the read, change and write are one step under STATE_LOCK, so a record written at the same time is not lost."""
    state_path = Path(state_path)
    with STATE_LOCK:
        whole = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
        if whole.get("workbooks", {}).get(_key(path)) == record:
            return                                                    # #181: the same record again (the same bytes): nothing to write
        whole.setdefault("workbooks", {})[_key(path)] = record
        state_path.parent.mkdir(parents=True, exist_ok=True)
        platform.write_file(state_path, json.dumps(whole, indent=1, ensure_ascii=False))


def has_unimported_edits(path, state_path):
    """True if the workbook exists and differs from both the last export and the last import
    (creaite excel_sync.has_unimported_edits). A workbook clark has no record of counts as edited: it is never overwritten."""
    h = file_hash(path)
    if h is None:
        return False
    state = load_state(state_path, path)
    return h not in (state.get("hash"), state.get("imported_hash"))


MAX_CREATED = 5000     # fingerprints of the rows an import created, kept per workbook; the oldest are dropped past this


def add_created(record, created):
    """Append `created` ([{"fp", "id", "step", "db_id", "at"}, ...]; the last three say which step made it) to a workbook record's "created" list, keeping the newest MAX_CREATED."""
    if created:
        keep = ("fp", "id", "step", "db_id", "at")
        record["created"] = (list(record.get("created") or []) + [{k: c[k] for k in keep if k in c} for c in created])[-MAX_CREATED:]
    return record


def created_for(state_path, path, project):
    """[{"fp", "id"}] the imports made that a read of `path` must not make again: this workbook's own record first, then those of
    other workbooks of the same project in the state file (a copy of a workbook, or the same one saved elsewhere, has no record of
    its own). clark addition (#38)."""
    state_path = Path(state_path)
    if not state_path.exists():
        return []
    books = json.loads(state_path.read_text(encoding="utf-8")).get("workbooks", {})
    own = books.get(_key(path), {})
    out = list(own.get("created") or [])
    for k, rec in books.items():
        if k != _key(path) and rec.get("project") == project:
            out += rec.get("created") or []
    return out


def mark_imported(path, state_path, created=None, project=None, step=None):
    """Record the workbook as it is now as imported, so the next export may overwrite it (what creaite's cmd_import does to
    the state at its end). For the import (#38). clark: `created` is the fingerprints of the rows this import made
    (record_created), written in the same save; `project` is kept if the record has none; `step` is the stamp of the step
    that applied the import ({"step", "db_id", "at"}), so "already imported" holds only while that step is in force."""
    with STATE_LOCK:                              # the record read and written back as one step
        record = load_state(state_path, path)
        if step:
            record["imported_step"] = step
        if project:
            record.setdefault("project", project)
        record["imported_hash"] = file_hash(path)
        add_created(record, created)
        save_state(state_path, path, record)


def record_created(path, state_path, created, project=None):
    """Record only the fingerprints of the rows an import created, without marking the file as imported (for a workbook saved
    again after it was read: its new rows must still not be made twice). clark addition (#38)."""
    if created:
        with STATE_LOCK:                          # the record read and written back as one step
            record = load_state(state_path, path)
            if project:
                record.setdefault("project", project)
            save_state(state_path, path, add_created(record, created))


def lock_file(path):
    """The file Excel keeps beside a workbook it has open (creaite excel_sync.LOCK)."""
    path = Path(path)
    return path.with_name("~$" + path.name)


def choices(table, field):
    """The drop-down values for a field, for this table first, then for the field anywhere (creaite excel_sync.choices)."""
    return CHOICES.get((table, field)) or CHOICES.get(field)


# The characters an .xlsx can't hold (XML 1.0 refuses them): the C0 controls but tab, line feed and carriage return, lone
# surrogates, and U+FFFE and U+FFFF (creaite excel_sync.NOT_XML).
NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


def cell_text(value):
    """A value as a cell may hold it: text without NOT_XML's characters (a vertical tab, form feed or separator becomes a
    space); anything else as it is (creaite excel_sync.cell_text). clark: a carriage return (or CR LF) becomes a line feed,
    which is what an XML file gives back anyway (XML 1.0 reads CR as LF), so the round trip is exact."""
    if not isinstance(value, str):
        return value
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    return NOT_XML.sub(lambda m: " " if m.group() in "\x0b\x0c\x1c\x1d\x1e" else "", value)


# A text cell that starts with one of these is run as a formula by Excel (a CSV/formula injection), so it is neutralised.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
APOSTROPHE = "'"


def escape_text(text):
    """Text as it goes in a cell, so no cell can be a formula. The one rule, reversed exactly by unescape_text: text that starts
    with = + - @ TAB CR, or with an apostrophe, gets one apostrophe in front (=1+1 becomes '=1+1; 'x becomes ''x). Other text
    is untouched. Call it after cell_text, which can remove characters from the front."""
    return APOSTROPHE + text if text.startswith(FORMULA_START + (APOSTROPHE,)) else text


def unescape_text(text):
    """Cell text to the stored text: removes exactly one leading apostrophe when what follows starts with = + - @ TAB CR or
    another apostrophe (the inverse of escape_text). Text typed in Excel, with no such apostrophe, comes back as typed."""
    if text.startswith(APOSTROPHE) and text[1:].startswith(FORMULA_START + (APOSTROPHE,)):
        return text[1:]
    return text


def to_cell(field, value):
    """A stored value as it goes in a cell: dates as real dates, flags as Yes or No, text without the characters a workbook
    can't hold (creaite excel_sync.to_cell). clark: no link fields, and text that could be a formula gets an apostrophe
    (escape_text). Numbers and dates are never touched."""
    if value is None:
        return None
    if field in DATES:
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return escape_text(cell_text(value))   # not a date: text as typed, which must not be a formula
    if field in YES_NO:
        return "Yes" if int(value) else "No"
    value = cell_text(value)
    return escape_text(value) if isinstance(value, str) else value


def from_cell(field, value):
    """Workbook value to stored value: dates YYYY-MM-DD, Yes/No to 1/0, blanks to None (creaite excel_sync.from_cell).
    clark: text has its escape apostrophe removed (unescape_text), and a formula is only ever its text, never run."""
    if value is None:
        return None
    if not isinstance(value, (str, int, float, bool, date, datetime)):
        value = str(getattr(value, "text", None) or value)     # an array formula object: its formula as text
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        if value.strip() == "":
            return None
        value = unescape_text(value)
        if field in DATES:
            try:
                return datetime.strptime(value.strip(), "%d-%m-%Y").date().isoformat()
            except ValueError:
                return value  # left as typed; the store's checks report it
    if field in YES_NO:
        if str(value).lower() in ("yes", "y", "1", "true"):
            return 1
        if str(value).lower() in ("no", "n", "0", "false"):
            return 0
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def add_list(ws, col, n_rows, values, allow_blank=True, source=None):
    """Add a drop-down to column `col`, rows 2 to `n_rows` (creaite excel_sync.add_list). The values are written in the rule
    itself, or `source` (a range such as Lists!$A$2:$A$9) is used when given. clark (#85c): the rule is a list, its error style
    is stop (anything else is refused), and the error is in plain words.

    The rule covers rows 2 to `n_rows` only, that is the last data row + EXTRA_ROWS, not the whole column: a row typed further
    down has no drop-down. A range source (Lists!...) is matched by Excel without regard to case, so "high" is accepted where
    "High" is listed (not verified in Excel); a literal list in the rule is case-sensitive. The import checks the value either way."""
    formula = source or '"' + ",".join(values) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, errorStyle="stop",
                        showErrorMessage=True, errorTitle=LIST_ERROR_TITLE, error=LIST_ERROR)
    dv.add(f"{col}2:{col}{n_rows}")
    ws.add_data_validation(dv)


def add_note(ws, col, n_rows, title, text):
    """A prompt that shows when a cell in column `col` is selected, and checks nothing (free text with a note)."""
    dv = DataValidation(showInputMessage=True, promptTitle=title, prompt=text)
    dv.add(f"{col}2:{col}{n_rows}")
    ws.add_data_validation(dv)


def id_prompt(ws, col, first, last, text):
    """An input message (no rule) on the ID cells `col`{first}:`col`{last}; it refuses nothing."""
    if last < first:
        return
    dv = DataValidation(allow_blank=True, showInputMessage=True, promptTitle=ID_PROMPT_TITLE, prompt=text)
    dv.add(f"{col}{first}:{col}{last}")
    ws.add_data_validation(dv)


def lock_column(ws, col, n_rows, label, message=None):
    """Refuse edits in column `col` (rows 2 to `n_rows`): `label` is set by clark (creaite excel_sync.lock_column).
    The message says what to do next."""
    dv = DataValidation(type="custom", formula1="FALSE", showErrorMessage=True, errorTitle="Not editable",
                        error=message or "This column is filled in by clark. Leave it as it is.")
    dv.add(f"{col}2:{col}{n_rows}")
    ws.add_data_validation(dv)


def _dxf(name):
    """The look of a colour name as a conditional format: the token's light fill and dark text."""
    fill, text = colours.hexes(name)
    return DifferentialStyle(font=Font(color=text), fill=PatternFill(fill_type="solid", start_color=fill, end_color=fill))


def add_colour_rules(ws, col, n_rows, values):
    """One rule per value: a cell equal to the word gets that word's colour, so a value changed in Excel recolours itself. The
    word is a fixed literal from excel_colours, never data."""
    for value, name in values.items():
        rule = Rule(type="cellIs", operator="equal", formula=['"' + value + '"'], dxf=_dxf(name))
        ws.conditional_formatting.add(f"{col}2:{col}{n_rows}", rule)


def late_formula(date_col, status_col, open_states):
    """The overdue rule for row 2 (it applies down the column): a real date before today on a row whose status is open."""
    is_open = ",".join(f'${status_col}2="{s}"' for s in sorted(open_states))
    return f"AND(ISNUMBER(${date_col}2),${date_col}2<TODAY(),OR({is_open}))"


def add_late_rule(ws, col, n_rows, date_col, status_col, open_states):
    """Red text on the date of an open item past its date. A formula rule, evaluated by Excel against today's date."""
    rule = FormulaRule(formula=[late_formula(date_col, status_col, open_states)], font=Font(color=colours.LATE_TEXT))
    ws.conditional_formatting.add(f"{col}2:{col}{n_rows}", rule)


def column_width(table, field):
    """A column's width by its kind: ID 10; the title 40 (48 when it is a long-text column); other long text 48; dates 12;
    flags and numbers 10; the Updated stamp 24 (it is text such as 2030-01-10 09:30:00.000); people, streams and labels 18;
    anything else 12."""
    if field == "id" or field in YES_NO or field in NUMBERS:
        return 10
    if field in WIDE:
        return 48
    if field == TITLE_OF.get(table):
        return 40
    if field in DATES:
        return 12
    if field == "updated_at":
        return 24
    return 18 if field in SHORT_NAMES else 12


def wraps(table, field):
    """Whether a column's text wraps: long text and the title column."""
    return field in WIDE or field == TITLE_OF.get(table)


def freeze_cell(table, cols):
    """Where the panes freeze: below the header row, and after the ID and the item's title (C2 on every sheet)."""
    fields = [f for _, f in cols]
    return f"{get_column_letter(fields.index(TITLE_OF[table]) + 2)}2"


def _ordered(table, rows):
    """A sheet's rows in order: by sort_order then ID for milestones and use cases, by ID for the rest (as the page lists
    them; creaite orders every sheet by its key)."""
    if table in ("milestones", "use_cases"):
        return sorted(rows, key=lambda r: (r.get("sort_order") or 0, r["id"]))
    return sorted(rows, key=lambda r: r["id"])


def _project_lists(project, rows_by_table):
    """The project's own lists for the drop-downs: {"workstream": [...names], "milestone": [...IDs], "use_case": [...IDs]}.
    Workstreams are those of the project (rows_by_table["workstreams"]), or, when that is not given, those the rows use.
    An empty list is left out (a drop-down with nothing in it would refuse everything). So a project with no milestones, or no
    use cases, has no list rule on those columns: they stay free text, and any ID typed there is checked on import instead."""
    streams = rows_by_table.get("workstreams")
    if streams is None:
        names = {r.get("workstream") for t in TABLES for r in rows_by_table.get(t, []) if r.get("workstream")}
    else:
        names = {r.get("name") for r in streams if r.get("name") and not r.get("archived")}
    out = {"workstream": sorted(names, key=lambda n: (n.lower(), n)),
           "milestone": sorted(r["id"] for r in rows_by_table.get("milestones", [])),
           "use_case": sorted(r["id"] for r in rows_by_table.get("use_cases", []))}
    return {k: [to_cell(k, v) for v in vs] for k, vs in out.items() if vs}


LIST_HEADS = {"workstream": "Workstreams", "milestone": "Milestones", "use_case": "Use cases"}


HISTORY_NAME = "History"
HISTORY_COLUMNS = (("When", "when_text", 22), ("Who", "who", 18), ("How", "how", 26), ("Item", "item", 34), ("Field", "field", 20),
                   ("From", "from", 44), ("To", "to", 44))
HISTORY_NOTE = "This sheet is read-only and for internal use. Do not send it to a client."
HISTORY_CELL_LIMIT = 5000                  # characters a history cell holds (design note, section 3)
HISTORY_FONT = "Arial"


def history_cell(value):
    """A history value as a cell holds it: text only, cut at HISTORY_CELL_LIMIT with an ellipsis."""
    if value is None:
        return None
    text = cell_text(str(value))      # protection is the string cell type (data_type "s"), as c27c96b: no visible apostrophe
    return text if len(text) <= HISTORY_CELL_LIMIT else text[:HISTORY_CELL_LIMIT - 1] + "\u2026"


def write_history(wb, view):
    """The History sheet (#123), after the data sheets: `view` is exactly what history.view_for_project returns, so the sheet and
    the Activity tab show the same rows. Row 1 is the internal note and the footnote, row 2 the header (frozen with it), then the
    rows newest first, then the closing line (or the empty words). Every cell is a text cell in Arial; the import never reads it."""
    ws = wb.create_sheet(HISTORY_NAME)

    def put(row, values, **kw):
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            if isinstance(v, str):
                c.data_type = "s"           # text is text: a value starting "=" is never a formula
            c.font = Font(name=HISTORY_FONT, **kw)
            c.alignment = Alignment(wrap_text=i in (6, 7), vertical="top")

    put(1, [history_cell(HISTORY_NOTE + " " + view["footnote"])], italic=True)
    ws.append([label for label, _, _ in HISTORY_COLUMNS])
    for i, (label, _, width) in enumerate(HISTORY_COLUMNS, start=1):
        head = ws.cell(row=2, column=i)
        head.font, head.fill = Font(name=HISTORY_FONT, bold=True, color=colours.HEADER_TEXT), HEADER
        head.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in view["rows"]:
        put(ws.max_row + 1, [history_cell(r[key]) for _, key, _ in HISTORY_COLUMNS])
    last = view["closing"] or (view["empty"] if not view["rows"] else "")
    if last:
        put(ws.max_row + 1, [history_cell(last)], italic=True)
    ws.freeze_panes = "A3"
    return ws


def write_lists(wb, lists):
    """The hidden Lists sheet: one column per project list, a header in row 1 (the ranges the drop-downs point at are those of
    `workbook`, in the same order). Hidden, not very hidden, so a person can unhide it; the import reads sheets by name and never opens it."""
    ws = wb.create_sheet(LISTS_NAME)
    for i, field in enumerate(lists, start=1):
        col = get_column_letter(i)
        head = ws.cell(row=1, column=i, value=LIST_HEADS[field])
        head.font, head.fill = HEADER_FONT, HEADER
        head.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[col].width = 24
        for n, v in enumerate(lists[field], start=2):
            c = ws.cell(row=n, column=i, value=v)
            c.data_type = "s"
    ws.sheet_state = "hidden"
    return ws


# ── Export ────────────────────────────────────────────────────────────────────

COVER_NAME = "Cover"
COVER_LABELS = {"internal": "Internal", "caution": "Caution", "return": "Sending changes back",
                "where": "Where to make changes", "ids": "IDs"}
COVER_FONT = "Arial"


def _cover_cell(value):
    """Cover text as a cell holds it: the same cell_text and escape_text as the data, so no cell can be a formula or hold a
    character XML can't."""
    return escape_text(cell_text(excel_cover.cut(value)))


def _cover_font(**kw):
    return Font(name=COVER_FONT, **kw)


def write_cover(wb, cover, mark=False):
    """The Cover sheet, first in the workbook and the one it opens on (#85b): one label column and one wide wrapped text column,
    no merged cells (Excel does not fit a merged cell's height to its text), frozen below the title band (row 1). Arial
    throughout. `cover` is excel_cover.lines(). Not a table: the import never reads it."""
    ws = wb.create_sheet(COVER_NAME, 0)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    navy, white = colours.TOKENS["brand-navy"], colours.TOKENS["white"]
    blue, grey_fill, grey_text = colours.TOKENS["brand-blue"], colours.TOKENS["grey-50"], colours.TOKENS["grey-700"]
    quiet = ("about", "exported", "caution", "return", "where", "ids", "internal")
    for line in cover:
        kind, text = line["kind"], line["text"]
        label = line.get("tag", "") if kind == "title" else COVER_LABELS.get(kind, "")
        if kind in ("heading", "about"):
            label, text = text, ""
        ws.append([_cover_cell(label), _cover_cell(text)])
        row = ws.max_row
        a, b = ws.cell(row=row, column=1), ws.cell(row=row, column=2)
        for c in (a, b):
            c.data_type = "s"           # text is text: openpyxl would make a string starting "=" a formula
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = _cover_font()
        if kind == "title":
            for c in (a, b):
                c.fill = PatternFill("solid", fgColor=navy)
                c.font = _cover_font(bold=True, size=14, color=white)
            if label:                   # the INTERNAL tag is red, as in Option A (white on the page's red, 5:1)
                a.fill = PatternFill("solid", fgColor=colours.TOKENS["status-red"])
        elif kind == "heading":
            a.font = _cover_font(bold=True, size=12, color=blue)
            a.border = Border(left=Side(style="thick", color=blue))
        elif kind == "health" and line.get("colour") in colours.COLOURS:     # the health pill, coloured as on the page
            fill, ink = colours.hexes(line["colour"])
            b.fill = PatternFill("solid", fgColor=fill)
            b.font = _cover_font(bold=True, color=ink)
        elif kind == "tile":
            b.fill = PatternFill("solid", fgColor=grey_fill)
        if kind in quiet:
            for c in (a, b):
                c.fill = PatternFill("solid", fgColor=grey_fill)
                c.font = _cover_font(bold=(kind == "about" and c is a), color=grey_text)
        if label and kind in ("caution", "return", "where", "ids", "internal"):
            a.font = _cover_font(bold=True, color=grey_text)
    ws.freeze_panes = "A2"
    if mark:
        ws.row_dimensions[1].height = excel_picture.COVER_ROW_POINTS      # room for the brand mark (#133), which excel_picture.embed adds
    wb.active = 0
    return ws


def workbook(project, rows_by_table, on=None, cover=None, mark=False, history=None):
    """The workbook for one project, from {table: [row, ...]} (creaite excel_sync.cmd_export, the sheet part). Rows are dicts
    with the table's columns; labelled rows may carry `labels` (comma-separated) and milestones `original_date`. A missing
    table is an empty sheet. `on` is the day Shown state is worked out for (today by default). `cover` (excel_cover.lines) adds
    the Cover as the first sheet (#85b); without it the workbook is as it was. `history` (history.view_for_project's result) adds
    the History sheet (#123) after the data sheets and before Lists."""
    lists = _project_lists(project, rows_by_table)
    ranges = {f: f"{LISTS_NAME}!${get_column_letter(i)}$2:${get_column_letter(i)}${len(vs) + 1}"
              for i, (f, vs) in enumerate(lists.items(), start=1)}
    wb = Workbook()
    wb.remove(wb.active)
    stamp = moment(on)
    props = wb.properties
    props.created = props.modified = stamp
    props.creator = props.lastModifiedBy = FIXED_NAME
    props.title = f"{project} work items"
    props.description = INTERNAL_NOTE
    for name, table, cols, locked in SHEETS:
        ws = wb.create_sheet(name)
        rows = _ordered(table, rows_by_table.get(table, []))
        ws.append([label for label, _ in cols])
        for r in rows:
            if table == "milestones":
                r = {**r, "shown_state": milestone_state(r, on), "original_date": r.get("original_date") or r.get("baseline_date")}
            ws.append([to_cell(f, r.get(f)) for _, f in cols])
            for c in ws[ws.max_row]:
                if isinstance(c.value, str):
                    c.data_type = "s"       # text is text: openpyxl would make a string starting "=" a formula, "#N/A" an error
        n_rows = len(rows) + 1 + EXTRA_ROWS
        for i, (label, field) in enumerate(cols, start=1):
            col = get_column_letter(i)
            head = ws.cell(row=1, column=i)
            head.font, head.fill = HEADER_FONT, HEADER
            head.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[col].width = column_width(table, field)
            for c in ws[f"{col}2:{col}{len(rows) + 1}"]:
                c[0].alignment = Alignment(wrap_text=wraps(table, field), vertical="top")
            if field in DATES:
                for c in ws[f"{col}2:{col}{n_rows}"]:
                    c[0].number_format = DATE_FORMAT
            if field == "id":
                id_prompt(ws, col, 2, len(rows) + 1, ID_PROMPT_EXISTING)
                id_prompt(ws, col, len(rows) + 2, n_rows, ID_PROMPT_NEW)
                head.comment = Comment(ID_NOTE, FIXED_NAME)
                head.comment.width, head.comment.height = 280, 60
            if field in locked:
                if field != "id":       # the ID is not refused by a rule: the import guards it (#129)
                    lock_column(ws, col, n_rows, label, "Priority follows Impact and Likelihood. Change those instead."
                                if (table, field) == ("raid", "priority") else None)
                for c in ws[f"{col}2:{col}{len(rows) + 1}"]:
                    c[0].fill = LOCKED
            blank = field not in YES_NO and (table, field) not in NOT_BLANK
            if field == "project":
                add_list(ws, col, n_rows, [project], allow_blank=False)
            elif field in ranges and field not in locked:
                add_list(ws, col, n_rows, None, allow_blank=True, source=ranges[field])   # the project's own list
            elif field in YES_NO:
                add_list(ws, col, n_rows, ["Yes", "No"], allow_blank=False)
            elif field not in locked and choices(table, field):
                add_list(ws, col, n_rows, choices(table, field), allow_blank=blank)
            elif field in PEOPLE:
                add_note(ws, col, n_rows, "Owner" if field == "owner" else "Made by",
                         "Type the person's name.")
            if (table, field) in colours.BY_FIELD:      # clark (#85c): colour follows the value, by conditional format
                add_colour_rules(ws, col, n_rows, colours.BY_FIELD[(table, field)])
            if (table, field) in colours.LATE:
                status_field, open_states = colours.LATE[(table, field)]
                add_late_rule(ws, col, n_rows, col, get_column_letter([f for _, f in cols].index(status_field) + 1),
                              open_states)
        if name == SHEETS[0][0]:
            note = Comment(INTERNAL_NOTE + "\n\n" + ID_NOTE, FIXED_NAME)    # one note per cell: both go on Milestones A1
            note.width, note.height = 320, 120
            ws["A1"].comment = note
        ws.freeze_panes = freeze_cell(table, cols)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    if history is not None:
        write_history(wb, history)
    if lists:
        write_lists(wb, lists)
    if cover:
        write_cover(wb, cover, mark)
    return wb


def fixed_zip(data, replace=None, add=None):
    """The same .xlsx with every member given one fixed date, in sorted order, deflated: openpyxl stamps each member with the
    time it was saved, so two saves in different seconds would differ. `replace` maps a member's name to the bytes to use, and `add` maps the name of a new member to its bytes."""
    replace = replace or {}
    add = add or {}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(set(src.namelist()) | set(add)):
            info = zipfile.ZipInfo(name, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, add[name] if name in add else replace[name] if name in replace else src.read(name))
    return out.getvalue()


def export(path, project, rows_by_table, state_path, on=None, cover=None, branding=None, warnings=None, history=None):
    """Write the workbook at `path` for `project` from {table: [row, ...]} and record what was written in `state_path`
    (creaite excel_sync.cmd_export). Returns {table: rows written}.

    Raises Refused, and leaves the file as it is, while the workbook is open in Excel or has edits not yet imported. clark: the
    workbook is saved with platform.write_file (a temporary file, then a safe replace with retry)."""
    path = Path(path)
    if lock_file(path).exists():
        raise Refused(f"{path.name} is open in Excel. Close it, then export again ({path}).")
    if has_unimported_edits(path, state_path):
        if load_state(state_path, path):
            raise Refused(f"{path.name} has edits that haven't been imported. Import it first; the workbook was not "
                          f"overwritten ({path}).")
        raise Refused(f"{path.name} already exists and wasn't written by an export from this store, so it was not "
                      f"overwritten ({path}). Import it, or choose another file name.")
    buf = io.BytesIO()
    logo = excel_picture.load(branding, warnings) if cover else None     # #133: the local brand logo, if there is one and it is safe
    wb = workbook(project, rows_by_table, on, cover, mark=bool(logo), history=history)
    wb.save(buf)
    # openpyxl sets "modified" to the real time as it saves, so core.xml is written again with the export's own stamp.
    wb.properties.modified = wb.properties.created
    replace, add = {"docProps/core.xml": tostring(wb.properties.to_tree())}, {}
    if logo:
        # The logo goes in after openpyxl has saved, by editing the zip's members (standard library only); the write below is
        # the atomic one. If the saved file is not laid out as expected the Cover goes without a mark.
        try:
            with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
                files = {n: z.read(n) for n in z.namelist()}
            more, add = excel_picture.embed(files, logo[0], COVER_NAME)
            replace.update(more)
        except excel_picture.PictureRefused as e:
            text = excel_picture.NOT_USED.format(why=e)         # the same plain sentence as a refused logo file
            excel_picture.LOG.warning("%s", text)
            if warnings is not None:
                warnings.append(text)
    data = fixed_zip(buf.getvalue(), replace, add)
    platform.write_file(path, data)
    snapshot = {t: {r["id"]: {k: v for k, v in r.items() if k not in DERIVED} for r in rows_by_table.get(t, [])}
                for t in TABLES}
    # clark: "rows" is what each row held when it was exported, {table: {id: row}}: the import compares a workbook with it to
    # tell the person's own edits from changes made in clark since. Left out (None) if it would be very large.
    if len(json.dumps(snapshot, default=str)) > MAX_SNAPSHOT_BYTES:
        snapshot_rows = None
    else:
        snapshot_rows = snapshot
    save_state(state_path, path, {"exported_at": moment(on).strftime("%Y-%m-%d %H:%M:%S"),
                                  "hash": hashlib.sha256(data).hexdigest(), "imported_hash": None, "project": project,
                                  "rows": snapshot_rows})
    return {t: len(snapshot[t]) for t in TABLES}


def export_name(project):
    """The file name clark gives a project's workbook, made safe (platform.safe_name)."""
    return platform.safe_name(f"{project} work items.xlsx")


def export_project(store, project, path=None, state_path=None, on=None, exported_by=None, owner=None, owner_email=None,
                   branding=None, warnings=None):
    """Export one project from the store (#31): each table's rows with Store.read(kind, project), archived rows left out (as
    creaite's hide_archived), labels from item_labels and each milestone's Original date from its history. `path` defaults
    to export_name() beside clark.db; the record goes in STATE_FILE beside clark.db. Returns the path written."""
    from clark import history
    folder = Path(store.path).parent
    path = Path(path) if path is not None else folder / export_name(project)
    state_path = Path(state_path) if state_path is not None else folder / STATE_FILE
    rows = {t: [r for r in store.read(t, project) if not r.get("archived")] for t in TABLES}
    rows["workstreams"] = [r for r in store.read("workstreams", project) if not r.get("archived")]   # for the drop-downs only
    ids = {r["id"] for t in LABELLED for r in rows[t]}
    have = {}
    for r in sorted(store.read("item_labels"), key=lambda r: r["label"].lower()):
        if r["item_id"] in ids:
            have.setdefault(r["item_id"], []).append(r["label"])
    for t in LABELLED:
        for r in rows[t]:
            r["labels"] = ", ".join(have.get(r["id"], [])) or None
    for r in rows["milestones"]:
        r["original_date"] = history.original_date(store.con, r["id"], r.get("baseline_date"))
    # The Cover is worked out from what the published page shows (excel_cover.from_store), so the two can't disagree; the person
    # named as exporting and the owner are the caller's to give (Settings are not read here).
    cover = excel_cover.from_store(store, project, exported_text(on), exported_by, owner, owner_email, INTERNAL_NOTE,
                                   day(on).isoformat())
    view = history.view_for_project(store.con, project, "all", history.ROW_LIMIT)       # the one source the Activity tab shares
    export(path, project, rows, state_path, on, cover, branding, warnings, history=view)
    return path


# ── Reading back ──────────────────────────────────────────────────────────────

def read_sheet(ws, table, cols):
    """Rows as {field: stored value}; header labels map to fields, so column order may change (creaite excel_sync.read_sheet).
    clark: worked-out columns are left out, and a blank cell in a NOT NULL text column reads as ""."""
    header = [c.value for c in ws[1]]
    by_label = {label: field for label, field in cols if field not in DERIVED}
    index = {i: by_label[h] for i, h in enumerate(header) if h in by_label}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = {index[i]: from_cell(index[i], v) for i, v in enumerate(row) if i in index}
        if any(v is not None for v in rec.values()):
            for f in rec:
                if rec[f] is None and (table, f) in EMPTY_TEXT:
                    rec[f] = ""
            rows.append(rec)
    return rows


def read_workbook(path):
    """{table: [row, ...]} as stored values, from an exported workbook (for the import, #38, and the round trip)."""
    check_workbook_size(path)
    wb = load_workbook(Path(path))
    return {table: read_sheet(wb[name], table, cols) if name in wb.sheetnames else [] for name, table, cols, _ in SHEETS}
